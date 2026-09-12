from datetime import datetime
from typing import List, Optional
from pydantic import BaseModel
from fastapi import APIRouter, HTTPException, Depends
from sqlalchemy.orm import Session
from sqlalchemy import func
from app.core.dependencies import get_db, require_teacher
from app.domain.models import (
    User, SchoolClass, Enrollment, Attendance, Grade, Homework, Submission, ClassSubject
)
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from fastapi.responses import StreamingResponse
from app.schemas.teacher import (
    CreateClassRequest, UpdateClassRequest, AddStudentRequest, UpdateStudentRequest,
    SaveAttendanceRequest, SaveGradesRequest, AssignSubjectTeacherRequest
)

router = APIRouter(prefix="/teacher", tags=["Teacher Management Flow"])

@router.get("/dashboard-stats")
def get_teacher_stats(teacher: User = Depends(require_teacher), db: Session = Depends(get_db)):
    teacher_id = teacher.id
    
    total_classes = db.query(SchoolClass).filter(SchoolClass.teacher_id == teacher_id).count()
    
    total_students = db.query(func.count(func.distinct(Enrollment.student_id))).join(
        SchoolClass, Enrollment.class_id == SchoolClass.id
    ).filter(SchoolClass.teacher_id == teacher_id).scalar() or 0
    
    total_homeworks = db.query(Homework).filter(Homework.teacher_id == teacher_id).count()
    
    pending_reviews = db.query(Submission).join(
        Homework, Submission.homework_id == Homework.id
    ).filter(
        Homework.teacher_id == teacher_id,
        Submission.score.is_(None)
    ).count()

    # Calculate overall attendance rate percentage
    today = datetime.utcnow().date()
    today_att = db.query(Attendance).join(
        SchoolClass, Attendance.class_id == SchoolClass.id
    ).filter(
        SchoolClass.teacher_id == teacher_id,
        Attendance.date == today
    ).all()
    
    if today_att and len(today_att) > 0:
        present_count = sum(1 for a in today_att if a.status == "present")
        attendance_rate = round((present_count / len(today_att)) * 100, 1)
    else:
        recent_att = db.query(Attendance).join(
            SchoolClass, Attendance.class_id == SchoolClass.id
        ).filter(
            SchoolClass.teacher_id == teacher_id
        ).limit(100).all()
        if recent_att and len(recent_att) > 0:
            present_count = sum(1 for a in recent_att if a.status == "present")
            attendance_rate = round((present_count / len(recent_att)) * 100, 1)
        else:
            attendance_rate = 96.0
    
    return {
        "total_classes": total_classes,
        "total_students": total_students,
        "attendance_rate": attendance_rate,
        "total_homeworks": total_homeworks,
        "pending_reviews": pending_reviews
    }


@router.get("/classes")
def list_classes(teacher: User = Depends(require_teacher), db: Session = Depends(get_db)):
    if teacher.role == "admin":
        classes = db.query(SchoolClass).order_by(SchoolClass.id.asc()).all()
    else:
        subject_class_ids = [cs.class_id for cs in db.query(ClassSubject.class_id).filter(ClassSubject.teacher_id == teacher.id).all()]
        classes = db.query(SchoolClass).filter(
            (SchoolClass.teacher_id == teacher.id) | (SchoolClass.id.in_(subject_class_ids))
        ).order_by(SchoolClass.id.asc()).all()
    
    results = []
    for c in classes:
        student_count = db.query(Enrollment).filter(Enrollment.class_id == c.id).count()
        assigned_subjects = [cs.subject_name for cs in db.query(ClassSubject).filter(ClassSubject.class_id == c.id, ClassSubject.teacher_id == teacher.id).all()]
        is_homeroom = (c.teacher_id == teacher.id or teacher.role == "admin")
        results.append({
            "id": c.id,
            "name": c.name,
            "grade_level": c.grade_level,
            "teacher_id": c.teacher_id,
            "teacher_name": c.teacher.name if c.teacher else "គ្មានគ្រូបង្រៀន",
            "academic_year": c.academic_year,
            "created_at": c.created_at,
            "student_count": student_count,
            "is_homeroom": is_homeroom,
            "assigned_subjects": assigned_subjects
        })
    return results

@router.post("/classes")
def create_class(payload: CreateClassRequest, teacher: User = Depends(require_teacher), db: Session = Depends(get_db)):
    target_teacher_id = teacher.id
    if teacher.role == "admin" and payload.teacher_id:
        target_teacher_id = payload.teacher_id

    new_class = SchoolClass(
        name=payload.name,
        grade_level=payload.grade_level,
        teacher_id=target_teacher_id,
        academic_year=payload.academic_year or "2025-2026"
    )
    db.add(new_class)
    db.commit()
    db.refresh(new_class)
    return {"success": True, "class_id": new_class.id, "name": new_class.name}

@router.put("/classes/{class_id}")
def update_class(
    class_id: int,
    payload: UpdateClassRequest,
    teacher: User = Depends(require_teacher),
    db: Session = Depends(get_db)
):
    school_class = db.query(SchoolClass).filter(SchoolClass.id == class_id).first()
    if not school_class:
        raise HTTPException(status_code=404, detail="រកមិនឃើញថ្នាក់រៀននេះទេ")
        
    if school_class.teacher_id != teacher.id and teacher.role != "admin":
        raise HTTPException(status_code=403, detail="អ្នកគ្មានសិទ្ធិកែប្រែថ្នាក់រៀននេះទេ")
        
    if payload.name is not None:
        school_class.name = payload.name.strip()
    if payload.grade_level is not None:
        school_class.grade_level = payload.grade_level.strip()
    if payload.academic_year is not None:
        school_class.academic_year = payload.academic_year.strip()
    if payload.teacher_id is not None and teacher.role == "admin":
        school_class.teacher_id = payload.teacher_id
        
    db.commit()
    db.refresh(school_class)
    return {"success": True, "message": "បានកែប្រែថ្នាក់រៀនដោយជោគជ័យ", "class": {
        "id": school_class.id,
        "name": school_class.name,
        "grade_level": school_class.grade_level,
        "academic_year": school_class.academic_year,
        "teacher_id": school_class.teacher_id
    }}

@router.delete("/classes/{class_id}")
def delete_class(
    class_id: int,
    teacher: User = Depends(require_teacher),
    db: Session = Depends(get_db)
):
    school_class = db.query(SchoolClass).filter(SchoolClass.id == class_id).first()
    if not school_class:
        raise HTTPException(status_code=404, detail="រកមិនឃើញថ្នាក់រៀននេះទេ")
        
    if school_class.teacher_id != teacher.id and teacher.role != "admin":
        raise HTTPException(status_code=403, detail="អ្នកគ្មានសិទ្ធិលុបថ្នាក់រៀននេះទេ")
        
    db.delete(school_class)
    db.commit()
    return {"success": True, "message": "បានលុបថ្នាក់រៀន និងទិន្នន័យពាក់ព័ន្ធដោយជោគជ័យ"}

@router.get("/classes/{class_id}/students")
def get_class_students(class_id: int, teacher: User = Depends(require_teacher), db: Session = Depends(get_db)):
    school_class = db.query(SchoolClass).filter(SchoolClass.id == class_id).first()
    if not school_class:
        raise HTTPException(status_code=404, detail="រកមិនឃើញថ្នាក់រៀននេះទេ")
    if school_class.teacher_id != teacher.id and teacher.role != "admin":
        raise HTTPException(status_code=403, detail="អ្នកគ្មានសិទ្ធិមើលសិស្សក្នុងថ្នាក់របស់គ្រូផ្សេងទេ")

    enrollments = db.query(Enrollment).join(
        User, Enrollment.student_id == User.id
    ).filter(
        Enrollment.class_id == class_id
    ).order_by(Enrollment.roll_no.asc(), User.name.asc()).all()
    
    students = []
    for e in enrollments:
        u = e.student
        students.append({
            "id": u.id,
            "name": u.name,
            "email": u.email,
            "student_code": u.student_code,
            "phone": u.phone,
            "avatar_url": u.avatar_url,
            "roll_no": e.roll_no
        })
    return students

@router.post("/classes/{class_id}/students")
def enroll_student(class_id: int, payload: AddStudentRequest, teacher: User = Depends(require_teacher), db: Session = Depends(get_db)):
    school_class = db.query(SchoolClass).filter(SchoolClass.id == class_id).first()
    if not school_class:
        raise HTTPException(status_code=404, detail="រកមិនឃើញថ្នាក់រៀននេះទេ")
    if school_class.teacher_id != teacher.id and teacher.role != "admin":
        raise HTTPException(status_code=403, detail="អ្នកគ្មានសិទ្ធិបញ្ចូលសិស្សក្នុងថ្នាក់របស់គ្រូផ្សេងទេ")

    email = payload.email.lower().strip()
    student = db.query(User).filter(func.lower(User.email) == email).first()
    
    if not student:
        user_count = db.query(User).count()
        student_code = payload.student_code or f"STU-{user_count + 1:03d}"
        student = User(
            email=email,
            name=payload.name,
            role="student",
            student_code=student_code,
            phone=payload.phone
        )
        db.add(student)
        db.commit()
        db.refresh(student)
        
    # Check already enrolled
    existing_enr = db.query(Enrollment).filter(
        Enrollment.class_id == class_id,
        Enrollment.student_id == student.id
    ).first()
    if existing_enr:
        raise HTTPException(status_code=400, detail="សិស្សនេះបានចុះឈ្មោះក្នុងថ្នាក់នេះរួចហើយ")
        
    max_roll = db.query(func.max(Enrollment.roll_no)).filter(Enrollment.class_id == class_id).scalar() or 0
    next_roll = max_roll + 1
    
    enr = Enrollment(class_id=class_id, student_id=student.id, roll_no=next_roll)
    db.add(enr)
    db.commit()
    
    return {"success": True, "student_id": student.id, "roll_no": next_roll}

@router.delete("/classes/{class_id}/students/{student_id}")
def unenroll_student(
    class_id: int,
    student_id: int,
    teacher: User = Depends(require_teacher),
    db: Session = Depends(get_db)
):
    school_class = db.query(SchoolClass).filter(SchoolClass.id == class_id).first()
    if not school_class:
        raise HTTPException(status_code=404, detail="រកមិនឃើញថ្នាក់រៀននេះទេ")
        
    if school_class.teacher_id != teacher.id and teacher.role != "admin":
        raise HTTPException(status_code=403, detail="អ្នកគ្មានសិទ្ធិគ្រប់គ្រងថ្នាក់រៀននេះទេ")
        
    enrollment = db.query(Enrollment).filter(
        Enrollment.class_id == class_id,
        Enrollment.student_id == student_id
    ).first()
    
    if not enrollment:
        raise HTTPException(status_code=404, detail="រកមិនឃើញសិស្សនេះក្នុងថ្នាក់ឡើយ")
        
    db.delete(enrollment)
    db.commit()
    return {"success": True, "message": "បានដកសិស្សចេញពីថ្នាក់រៀនដោយជោគជ័យ"}

@router.put("/students/{student_id}")
def update_student(
    student_id: int,
    payload: UpdateStudentRequest,
    teacher: User = Depends(require_teacher),
    db: Session = Depends(get_db)
):
    student = db.query(User).filter(User.id == student_id, User.role == "student").first()
    if not student:
        raise HTTPException(status_code=404, detail="រកមិនឃើញសិស្សនេះទេ")
        
    if payload.name is not None:
        student.name = payload.name.strip()
    if payload.phone is not None:
        student.phone = payload.phone.strip()
    if payload.student_code is not None:
        # Check uniqueness if changed
        new_code = payload.student_code.strip()
        existing = db.query(User).filter(User.student_code == new_code, User.id != student_id).first()
        if existing:
            raise HTTPException(status_code=400, detail="អត្តលេខសិស្សនេះមានរួចហើយ (Student code duplicate)")
        student.student_code = new_code
    if payload.email is not None:
        new_email = payload.email.strip().lower()
        existing_email = db.query(User).filter(func.lower(User.email) == new_email, User.id != student_id).first()
        if existing_email:
            raise HTTPException(status_code=400, detail="អ៊ីមែលនេះមានរួចហើយ (Email duplicate)")
        student.email = new_email
        
    db.commit()
    db.refresh(student)
    return {
        "success": True,
        "message": "បានកែប្រែព័ត៌មានសិស្សដោយជោគជ័យ",
        "student": {
            "id": student.id,
            "name": student.name,
            "email": student.email,
            "student_code": student.student_code,
            "phone": student.phone
        }
    }

class TransferStudentRequest(BaseModel):
    current_class_id: int
    target_class_id: int

@router.get("/all-students")
def get_all_students(teacher: User = Depends(require_teacher), db: Session = Depends(get_db)):
    """
    Returns ONLY role='student' users in the teacher's classes.
    Teachers and admins are strictly excluded from the response.
    """
    if teacher.role == "admin":
        # Admin: all enrollments, but only for users with role='student'
        enrollments = (
            db.query(Enrollment)
            .join(User, Enrollment.student_id == User.id)
            .filter(User.role == "student")          # ← strict role filter
            .all()
        )
    else:
        teacher_classes = db.query(SchoolClass).filter(SchoolClass.teacher_id == teacher.id).all()
        class_ids = [c.id for c in teacher_classes]
        if not class_ids:
            return []
        enrollments = (
            db.query(Enrollment)
            .join(User, Enrollment.student_id == User.id)
            .filter(
                Enrollment.class_id.in_(class_ids),
                User.role == "student"               # ← strict role filter
            )
            .all()
        )

    results = []
    seen_ids: set[int] = set()

    for enr in enrollments:
        stu = enr.student
        # Double-check: skip non-students even if somehow joined
        if not stu or stu.role != "student" or stu.id in seen_ids:
            continue
        cls = enr.school_class
        results.append({
            "id": stu.id,
            "student_id": stu.student_code or f"STU-{stu.id:03d}",
            "student_code": stu.student_code or f"STU-{stu.id:03d}",
            "full_name": stu.name,
            "name": stu.name,
            "email": stu.email,
            "phone": stu.phone or "-",
            "phone_number": stu.phone or "-",
            "gender": "male",
            "status": "active",
            "role": "student",
            "class_id": cls.id if cls else None,
            "enrolled_class": cls.name if cls else "-",
            "class_name": cls.name if cls else "-",
            "grade_level": cls.grade_level if cls else "-",
            "roll_no": enr.roll_no,
            "avatar_url": stu.avatar_url,
        })
        seen_ids.add(stu.id)

    # Admin also sees unassigned students (role='student', not in any class)
    if teacher.role == "admin":
        unassigned = (
            db.query(User)
            .filter(
                User.role == "student",              # ← strict role filter
                ~User.id.in_(seen_ids) if seen_ids else True
            )
            .all()
        )
        for stu in unassigned:
            results.append({
                "id": stu.id,
                "student_id": stu.student_code or f"STU-{stu.id:03d}",
                "student_code": stu.student_code or f"STU-{stu.id:03d}",
                "full_name": stu.name,
                "name": stu.name,
                "email": stu.email,
                "phone": stu.phone or "-",
                "phone_number": stu.phone or "-",
                "gender": "male",
                "status": "active",
                "role": "student",
                "class_id": None,
                "enrolled_class": "មិនទាន់មានថ្នាក់ (Unassigned)",
                "class_name": "មិនទាន់មានថ្នាក់",
                "grade_level": "-",
                "roll_no": None,
                "avatar_url": stu.avatar_url,
            })

    return results

@router.post("/students/{student_id}/transfer")
def transfer_student(
    student_id: int,
    payload: TransferStudentRequest,
    teacher: User = Depends(require_teacher),
    db: Session = Depends(get_db)
):
    target_class = db.query(SchoolClass).filter(SchoolClass.id == payload.target_class_id).first()
    if not target_class:
        raise HTTPException(status_code=404, detail="រកមិនឃើញថ្នាក់គោលដៅទេ")
        
    current_class = db.query(SchoolClass).filter(SchoolClass.id == payload.current_class_id).first()
    if not current_class:
        raise HTTPException(status_code=404, detail="រកមិនឃើញថ្នាក់បច្ចុប្បន្នទេ")

    if teacher.role != "admin" and (current_class.teacher_id != teacher.id or target_class.teacher_id != teacher.id):
        raise HTTPException(status_code=403, detail="អ្នកអាចផ្ទេរសិស្សបានតែរវាងថ្នាក់ដែលអ្នកទទួលបន្ទុកប៉ុណ្ណោះ")
        
    enrollment = db.query(Enrollment).filter(
        Enrollment.class_id == payload.current_class_id,
        Enrollment.student_id == student_id
    ).first()
    
    if not enrollment:
        raise HTTPException(status_code=404, detail="សិស្សមិនមានក្នុងថ្នាក់បច្ចុប្បន្នឡើយ")
        
    enrollment.class_id = payload.target_class_id
    db.commit()
    return {"success": True, "message": "បានផ្ទេរសិស្សទៅកាន់ថ្នាក់ថ្មីដោយជោគជ័យ"}

@router.delete("/students/{student_id}")
def delete_student_permanently(
    student_id: int,
    teacher: User = Depends(require_teacher),
    db: Session = Depends(get_db)
):
    if teacher.role != "admin":
        raise HTTPException(status_code=403, detail="មានតែ Admin ប៉ុណ្ណោះដែលអាចលុបគណនីសិស្សជាស្ថាពរបាន")

    student = db.query(User).filter(User.id == student_id, User.role == "student").first()
    if not student:
        raise HTTPException(status_code=404, detail="រកមិនឃើញសិស្សនេះទេ")
        
    db.delete(student)
    db.commit()
    return {"success": True, "message": "បានលុបព័ត៌មានសិស្សចេញពីប្រព័ន្ធដោយជោគជ័យ"}

@router.get("/attendance")
def get_attendance(
    class_id: int,
    date_str: str,
    subject: Optional[str] = "ទូទៅ",
    teacher: User = Depends(require_teacher),
    db: Session = Depends(get_db)
):
    school_class = db.query(SchoolClass).filter(SchoolClass.id == class_id).first()
    if not school_class:
        raise HTTPException(status_code=404, detail="រកមិនឃើញថ្នាក់រៀននេះទេ")
        
    is_subject_teacher = db.query(ClassSubject).filter(
        ClassSubject.class_id == class_id,
        ClassSubject.teacher_id == teacher.id
    ).first() is not None

    if school_class.teacher_id != teacher.id and not is_subject_teacher and teacher.role != "admin":
        raise HTTPException(status_code=403, detail="អ្នកគ្មានសិទ្ធិមើលវត្តមានថ្នាក់របស់គ្រូផ្សេងទេ")

    target_subject = subject.strip() if subject and subject.strip() else "ទូទៅ"
    enrollments = db.query(Enrollment).filter(Enrollment.class_id == class_id).order_by(Enrollment.roll_no.asc()).all()
    parsed_date = datetime.strptime(date_str, "%Y-%m-%d").date()
    
    records = []
    for e in enrollments:
        u = e.student
        att = db.query(Attendance).filter(
            Attendance.class_id == class_id,
            Attendance.student_id == u.id,
            Attendance.subject == target_subject,
            Attendance.date == parsed_date
        ).first()

        # Cumulative attendance stats for this student in this class strictly for this subject
        att_stats = db.query(Attendance.status, func.count(Attendance.id)).filter(
            Attendance.class_id == class_id,
            Attendance.student_id == u.id,
            Attendance.subject == target_subject
        ).group_by(Attendance.status).all()
        stat_dict = {st: cnt for st, cnt in att_stats}
        present_cnt = stat_dict.get("present", 0)
        late_cnt = stat_dict.get("late", 0)
        perm_cnt = stat_dict.get("permission", 0)
        absent_cnt = stat_dict.get("absent", 0)
        tot_days = present_cnt + late_cnt + perm_cnt + absent_cnt
        rate = round(((present_cnt + late_cnt * 0.8 + perm_cnt * 0.5) / tot_days * 100) if tot_days > 0 else 100.0, 0)
        
        records.append({
            "student_id": u.id,
            "name": u.name,
            "student_code": u.student_code,
            "roll_no": e.roll_no,
            "subject": target_subject,
            "status": att.status if att else "present",
            "notes": att.notes if att else None,
            "present_count": present_cnt,
            "attended_count": present_cnt,
            "late_count": late_cnt,
            "permission_count": perm_cnt,
            "absent_count": absent_cnt,
            "total_days": tot_days,
            "attendance_rate": int(rate)
        })
    return records

@router.post("/attendance")
def save_attendance(payload: SaveAttendanceRequest, teacher: User = Depends(require_teacher), db: Session = Depends(get_db)):
    school_class = db.query(SchoolClass).filter(SchoolClass.id == payload.class_id).first()
    if not school_class:
        raise HTTPException(status_code=404, detail="រកមិនឃើញថ្នាក់រៀននេះទេ")

    is_subject_teacher = db.query(ClassSubject).filter(
        ClassSubject.class_id == payload.class_id,
        ClassSubject.teacher_id == teacher.id
    ).first() is not None

    if school_class.teacher_id != teacher.id and not is_subject_teacher and teacher.role != "admin":
        raise HTTPException(status_code=403, detail="អ្នកគ្មានសិទ្ធិកត់ត្រាវត្តមានក្នុងថ្នាក់របស់គ្រូផ្សេងទេ")

    parsed_date = datetime.strptime(payload.date, "%Y-%m-%d").date()
    target_subject = payload.subject.strip() if payload.subject and payload.subject.strip() else "ទូទៅ"
    
    for item in payload.records:
        att = db.query(Attendance).filter(
            Attendance.class_id == payload.class_id,
            Attendance.student_id == item.student_id,
            Attendance.subject == target_subject,
            Attendance.date == parsed_date
        ).first()
        
        if att:
            att.status = item.status
            att.notes = item.notes
            att.teacher_id = teacher.id
        else:
            new_att = Attendance(
                class_id=payload.class_id,
                student_id=item.student_id,
                teacher_id=teacher.id,
                subject=target_subject,
                date=parsed_date,
                status=item.status,
                notes=item.notes
            )
            db.add(new_att)
            
    db.commit()
    return {"success": True, "count": len(payload.records), "date": payload.date, "subject": target_subject}

@router.get("/classes/{class_id}/subjects")
def get_class_subjects(class_id: int, teacher: User = Depends(require_teacher), db: Session = Depends(get_db)):
    subjects = db.query(ClassSubject).filter(ClassSubject.class_id == class_id).all()
    return [
        {
            "id": s.id,
            "class_id": s.class_id,
            "teacher_id": s.teacher_id,
            "teacher_name": s.teacher.name if s.teacher else "គ្រូបង្រៀន",
            "teacher_email": s.teacher.email if s.teacher else "",
            "subject_name": s.subject_name,
            "created_at": s.created_at
        }
        for s in subjects
    ]

@router.post("/classes/{class_id}/subjects")
def assign_class_subject(
    class_id: int,
    payload: AssignSubjectTeacherRequest,
    teacher: User = Depends(require_teacher),
    db: Session = Depends(get_db)
):
    school_class = db.query(SchoolClass).filter(SchoolClass.id == class_id).first()
    if not school_class:
        raise HTTPException(status_code=404, detail="រកមិនឃើញថ្នាក់រៀននេះទេ")
    if school_class.teacher_id != teacher.id and teacher.role != "admin":
        raise HTTPException(status_code=403, detail="មានតែ Admin ឬគ្រូប្រចាំថ្នាក់ទេដែលអាចចាត់តាំងគ្រូតាមមុខវិជ្ជា")

    target_teacher = db.query(User).filter(User.id == payload.teacher_id).first()
    if not target_teacher or target_teacher.role not in ["teacher", "admin"]:
        raise HTTPException(status_code=400, detail="សូមជ្រើសរើសគណនីគ្រូបង្រៀនត្រឹមត្រូវ")

    subj_name = payload.subject_name.strip()
    existing = db.query(ClassSubject).filter(
        ClassSubject.class_id == class_id,
        ClassSubject.subject_name == subj_name
    ).first()

    if existing:
        existing.teacher_id = target_teacher.id
        db.commit()
        return {"success": True, "message": f"បានធ្វើបច្ចុប្បន្នភាពគ្រូបង្រៀនមុខវិជ្ជា {subj_name} ទៅកាន់ {target_teacher.name}"}
    
    new_sub = ClassSubject(
        class_id=class_id,
        teacher_id=target_teacher.id,
        subject_name=subj_name
    )
    db.add(new_sub)
    db.commit()
    db.refresh(new_sub)
    return {"success": True, "message": f"បានចាត់តាំងលោកគ្រូ/អ្នកគ្រូ {target_teacher.name} បង្រៀនមុខវិជ្ជា {subj_name} ដោយជោគជ័យ"}

@router.delete("/classes/{class_id}/subjects/{subject_id}")
def delete_class_subject(
    class_id: int,
    subject_id: int,
    teacher: User = Depends(require_teacher),
    db: Session = Depends(get_db)
):
    school_class = db.query(SchoolClass).filter(SchoolClass.id == class_id).first()
    if not school_class:
        raise HTTPException(status_code=404, detail="រកមិនឃើញថ្នាក់រៀននេះទេ")
    if school_class.teacher_id != teacher.id and teacher.role != "admin":
        raise HTTPException(status_code=403, detail="អ្នកគ្មានសិទ្ធិលុបការចាត់តាំងនេះទេ")

    sub = db.query(ClassSubject).filter(ClassSubject.id == subject_id, ClassSubject.class_id == class_id).first()
    if not sub:
        raise HTTPException(status_code=404, detail="រកមិនឃើញមុខវិជ្ជានេះទេ")
    db.delete(sub)
    db.commit()
    return {"success": True, "message": "បានដកការចាត់តាំងមុខវិជ្ជាដោយជោគជ័យ"}

@router.get("/grades")
def get_grades_matrix(
    class_id: int,
    subject: Optional[str] = None,
    teacher: User = Depends(require_teacher),
    db: Session = Depends(get_db)
):
    school_class = db.query(SchoolClass).filter(SchoolClass.id == class_id).first()
    if not school_class:
        raise HTTPException(status_code=404, detail="រកមិនឃើញថ្នាក់រៀននេះទេ")

    is_homeroom = (school_class.teacher_id == teacher.id) or (teacher.role == "admin")

    # Get subjects this teacher is assigned to in this class
    assigned_subjects = [
        cs.subject_name for cs in db.query(ClassSubject).filter(
            ClassSubject.class_id == class_id,
            ClassSubject.teacher_id == teacher.id
        ).all()
    ]
    is_subject_teacher = len(assigned_subjects) > 0

    if not is_homeroom and not is_subject_teacher:
        raise HTTPException(status_code=403, detail="អ្នកគ្មានសិទ្ធិមើលពិន្ទុថ្នាក់របស់គ្រូផ្សេងទេ")

    enrollments = db.query(Enrollment).filter(Enrollment.class_id == class_id).order_by(Enrollment.roll_no.asc()).all()

    # Determine which subjects to show based on role
    grades_query = db.query(Grade).filter(Grade.class_id == class_id)
    if not is_homeroom and is_subject_teacher:
        # Subject teacher: only see their assigned subjects
        if subject and subject.strip():
            # Filter requested subject — must be in their assigned list
            subj = subject.strip()
            if subj not in assigned_subjects:
                raise HTTPException(status_code=403, detail="អ្នកគ្មានសិទ្ធិមើលពិន្ទុមុខវិជ្ជានេះទេ")
            grades_query = grades_query.filter(Grade.subject == subj)
        else:
            grades_query = grades_query.filter(Grade.subject.in_(assigned_subjects))
    elif subject and subject.strip():
        grades_query = grades_query.filter(Grade.subject == subject.strip())

    all_grades = grades_query.order_by(Grade.date.desc()).all()

    student_scores = {}
    for e in enrollments:
        s = e.student
        s_id = s.id
        s_grades = [
            {
                "student_id": g.student_id,
                "subject": g.subject,
                "exam_type": g.exam_type,
                "score": float(g.score),
                "max_score": float(g.max_score),
                "date": str(g.date)
            }
            for g in all_grades if g.student_id == s_id
        ]

        if s_grades:
            avg = float(sum(g["score"] for g in s_grades) / len(s_grades))
        else:
            avg = 0.0

        student_scores[s_id] = {
            "student_id": s.id,
            "name": s.name,
            "student_code": s.student_code,
            "roll_no": e.roll_no,
            "grades": s_grades,
            "average": round(avg, 2),
            "letter_grade": "A" if avg >= 90 else ("B" if avg >= 80 else ("C" if avg >= 70 else ("D" if avg >= 60 else ("E" if avg >= 50 else "F")))),
            "assigned_subjects": assigned_subjects if is_subject_teacher and not is_homeroom else None
        }

    ranked = sorted(student_scores.values(), key=lambda x: x["average"], reverse=True)
    for rank, item in enumerate(ranked, start=1):
        item["rank"] = rank if item["average"] > 0 else "-"

    return sorted(ranked, key=lambda x: x["roll_no"] if x["roll_no"] is not None else 0)


@router.get("/rankings")
def get_rankings(
    scope: str,  # 'class', 'subject', 'school'
    class_id: Optional[int] = None,
    subject: Optional[str] = None,
    grade_level: Optional[str] = None,
    teacher: User = Depends(require_teacher),
    db: Session = Depends(get_db)
):
    """Rankings endpoint supporting three scopes:
    - scope=class: rank students by overall average within a class
    - scope=subject: rank students by a specific subject within a class
    - scope=school: rank all students of a given grade_level across all classes
    """
    def letter(avg):
        if avg >= 90: return "A"
        elif avg >= 80: return "B"
        elif avg >= 70: return "C"
        elif avg >= 60: return "D"
        elif avg >= 50: return "E"
        return "F"

    if scope == "class":
        if not class_id:
            raise HTTPException(status_code=400, detail="ត្រូវការ class_id")
        school_class = db.query(SchoolClass).filter(SchoolClass.id == class_id).first()
        if not school_class:
            raise HTTPException(status_code=404, detail="រកមិនឃើញថ្នាក់រៀននេះទេ")

        # Access check
        is_homeroom = (school_class.teacher_id == teacher.id) or (teacher.role == "admin")
        assigned_subjects = [cs.subject_name for cs in db.query(ClassSubject).filter(
            ClassSubject.class_id == class_id, ClassSubject.teacher_id == teacher.id
        ).all()]
        if not is_homeroom and not assigned_subjects:
            raise HTTPException(status_code=403, detail="អ្នកគ្មានសិទ្ធិមើលចំណាត់ថ្នាក់ក្នុងថ្នាក់នេះទេ")

        enrollments = db.query(Enrollment).filter(Enrollment.class_id == class_id).all()
        grades_query = db.query(Grade).filter(Grade.class_id == class_id)
        if not is_homeroom:
            grades_query = grades_query.filter(Grade.subject.in_(assigned_subjects))
        all_grades = grades_query.all()

        results = []
        for e in enrollments:
            s_grades = [g for g in all_grades if g.student_id == e.student_id]
            avg = float(sum(float(g.score) for g in s_grades) / len(s_grades)) if s_grades else 0.0
            results.append({
                "student_id": e.student_id,
                "name": e.student.name,
                "student_code": e.student.student_code,
                "roll_no": e.roll_no,
                "class_name": school_class.name,
                "grade_level": school_class.grade_level,
                "average": round(avg, 2),
                "letter_grade": letter(avg),
                "subject_scores": [
                    {"subject": g.subject, "score": float(g.score)}
                    for g in s_grades
                ]
            })

        results.sort(key=lambda x: x["average"], reverse=True)
        for i, r in enumerate(results, 1):
            r["rank"] = i if r["average"] > 0 else "-"
        return {"scope": "class", "class_name": school_class.name, "grade_level": school_class.grade_level, "results": results}

    elif scope == "subject":
        if not class_id or not subject:
            raise HTTPException(status_code=400, detail="ត្រូវការ class_id និង subject")
        school_class = db.query(SchoolClass).filter(SchoolClass.id == class_id).first()
        if not school_class:
            raise HTTPException(status_code=404, detail="រកមិនឃើញថ្នាក់រៀននេះទេ")

        is_homeroom = (school_class.teacher_id == teacher.id) or (teacher.role == "admin")
        assigned_subjects = [cs.subject_name for cs in db.query(ClassSubject).filter(
            ClassSubject.class_id == class_id, ClassSubject.teacher_id == teacher.id
        ).all()]
        subj = subject.strip()
        if not is_homeroom and subj not in assigned_subjects:
            raise HTTPException(status_code=403, detail="អ្នកគ្មានសិទ្ធិមើលចំណាត់ថ្នាក់មុខវិជ្ជានេះទេ")

        enrollments = db.query(Enrollment).filter(Enrollment.class_id == class_id).all()
        grades = db.query(Grade).filter(Grade.class_id == class_id, Grade.subject == subj).all()
        grade_map = {g.student_id: float(g.score) for g in grades}

        results = []
        for e in enrollments:
            score = grade_map.get(e.student_id, None)
            results.append({
                "student_id": e.student_id,
                "name": e.student.name,
                "student_code": e.student.student_code,
                "roll_no": e.roll_no,
                "class_name": school_class.name,
                "grade_level": school_class.grade_level,
                "score": round(score, 2) if score is not None else None,
                "letter_grade": letter(score) if score is not None else "-"
            })

        results.sort(key=lambda x: (x["score"] is None, -(x["score"] or 0)))
        for i, r in enumerate(results, 1):
            r["rank"] = i if r["score"] is not None else "-"
        return {"scope": "subject", "subject": subj, "class_name": school_class.name, "grade_level": school_class.grade_level, "results": results}

    elif scope == "school":
        # School-wide ranking across all classes of a given grade_level
        # Only accessible to homeroom/admin teachers or subject teachers in that grade
        if not grade_level:
            raise HTTPException(status_code=400, detail="ត្រូវការ grade_level")

        if teacher.role == "admin":
            classes = db.query(SchoolClass).filter(SchoolClass.grade_level == grade_level).all()
        else:
            # Teachers can see school-wide ranking only for grade levels they teach
            teacher_class_ids = [c.id for c in db.query(SchoolClass).filter(SchoolClass.teacher_id == teacher.id).all()]
            subject_class_ids = [cs.class_id for cs in db.query(ClassSubject).filter(ClassSubject.teacher_id == teacher.id).all()]
            all_accessible_class_ids = list(set(teacher_class_ids + subject_class_ids))
            classes = db.query(SchoolClass).filter(
                SchoolClass.grade_level == grade_level,
                SchoolClass.id.in_(all_accessible_class_ids)
            ).all()

        if not classes:
            raise HTTPException(status_code=404, detail="រកមិនឃើញថ្នាក់នៅថ្នាក់ទីនេះទេ")

        class_ids = [c.id for c in classes]
        class_map = {c.id: c for c in classes}

        all_enrollments = db.query(Enrollment).filter(Enrollment.class_id.in_(class_ids)).all()
        all_grades_q = db.query(Grade).filter(Grade.class_id.in_(class_ids))
        if subject and subject.strip():
            all_grades_q = all_grades_q.filter(Grade.subject == subject.strip())
        all_grades = all_grades_q.all()

        student_grade_map = {}
        for g in all_grades:
            if g.student_id not in student_grade_map:
                student_grade_map[g.student_id] = []
            student_grade_map[g.student_id].append(float(g.score))

        results = []
        seen_students = set()
        for e in all_enrollments:
            if e.student_id in seen_students:
                continue
            seen_students.add(e.student_id)
            scores = student_grade_map.get(e.student_id, [])
            avg = float(sum(scores) / len(scores)) if scores else 0.0
            cls = class_map.get(e.class_id)
            results.append({
                "student_id": e.student_id,
                "name": e.student.name,
                "student_code": e.student.student_code,
                "roll_no": e.roll_no,
                "class_name": cls.name if cls else "",
                "grade_level": grade_level,
                "average": round(avg, 2),
                "letter_grade": letter(avg)
            })

        results.sort(key=lambda x: x["average"], reverse=True)
        for i, r in enumerate(results, 1):
            r["rank"] = i if r["average"] > 0 else "-"
        return {"scope": "school", "grade_level": grade_level, "subject": subject or "ទាំងអស់", "results": results}

    else:
        raise HTTPException(status_code=400, detail="scope ត្រូវជា class, subject, ឬ school")

@router.post("/grades")
def save_grades(payload: SaveGradesRequest, teacher: User = Depends(require_teacher), db: Session = Depends(get_db)):
    school_class = db.query(SchoolClass).filter(SchoolClass.id == payload.class_id).first()
    if not school_class:
        raise HTTPException(status_code=404, detail="រកមិនឃើញថ្នាក់រៀននេះទេ")

    is_subject_teacher = db.query(ClassSubject).filter(
        ClassSubject.class_id == payload.class_id,
        ClassSubject.teacher_id == teacher.id
    ).first() is not None

    if school_class.teacher_id != teacher.id and not is_subject_teacher and teacher.role != "admin":
        raise HTTPException(status_code=403, detail="អ្នកគ្មានសិទ្ធិកត់ត្រាពិន្ទុក្នុងថ្នាក់របស់គ្រូផ្សេងទេ")

    parsed_date = datetime.strptime(payload.date, "%Y-%m-%d").date()
    
    for item in payload.records:
        existing = db.query(Grade).filter(
            Grade.class_id == payload.class_id,
            Grade.student_id == item.student_id,
            Grade.subject == payload.subject,
            Grade.exam_type == payload.exam_type
        ).first()
        
        if existing:
            existing.score = item.score
            existing.max_score = item.max_score or 100.0
            existing.date = parsed_date
        else:
            new_grade = Grade(
                class_id=payload.class_id,
                student_id=item.student_id,
                subject=payload.subject,
                exam_type=payload.exam_type,
                score=item.score,
                max_score=item.max_score or 100.0,
                date=parsed_date
            )
            db.add(new_grade)
            
    db.commit()
    return {"success": True, "saved_records": len(payload.records)}

@router.get("/grades/export")
def export_grades_excel(
    class_id: int,
    teacher: User = Depends(require_teacher),
    db: Session = Depends(get_db)
):
    school_class = db.query(SchoolClass).filter(SchoolClass.id == class_id).first()
    if not school_class:
        raise HTTPException(status_code=404, detail="រកមិនឃើញថ្នាក់រៀននេះទេ")
        
    enrollments = db.query(Enrollment).filter(Enrollment.class_id == class_id).order_by(Enrollment.roll_no.asc()).all()
    all_grades = db.query(Grade).filter(Grade.class_id == class_id).all()
    
    # Extract distinct subjects
    subjects = sorted(list(set(g.subject for g in all_grades)))
    
    # Calculate scores per student
    student_scores = []
    for e in enrollments:
        s = e.student
        s_grades = {g.subject: float(g.score) for g in all_grades if g.student_id == s.id}
        total = sum(s_grades.values())
        avg = (total / len(s_grades)) if s_grades else 0.0
        
        student_scores.append({
            "roll_no": e.roll_no,
            "name": s.name,
            "student_code": s.student_code or "",
            "scores": s_grades,
            "total": round(total, 2),
            "average": round(avg, 2),
            "letter": "A" if avg >= 90 else ("B" if avg >= 80 else ("C" if avg >= 70 else ("D" if avg >= 60 else ("E" if avg >= 50 else "F"))))
        })
        
    # Rank students by average
    ranked = sorted(student_scores, key=lambda x: x["average"], reverse=True)
    for idx, item in enumerate(ranked, start=1):
        item["rank"] = idx if item["average"] > 0 else "-"
        
    student_scores.sort(key=lambda x: x["roll_no"] if x["roll_no"] is not None else 0)
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Grades Report"
    
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center")
    
    # Headers
    headers = ["ល.រ (Roll)", "អត្តលេខ (Code)", "គោត្តនាម-នាម (Name)"] + subjects + ["សរុប (Total)", "មធ្យមភាគ (Avg)", "និទ្ទេស (Grade)", "ចំណាត់ថ្នាក់ (Rank)"]
    ws.append(headers)
    
    for col_num, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
        
    for item in student_scores:
        row = [item["roll_no"], item["student_code"], item["name"]]
        for subj in subjects:
            row.append(item["scores"].get(subj, "-"))
        row.extend([item["total"], item["average"], item["letter"], item["rank"]])
        ws.append(row)
        
    # Auto adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
        
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"grades_{school_class.name.replace(' ', '_')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )

@router.get("/attendance/export")
def export_attendance_excel(
    class_id: int,
    month: Optional[str] = None, # format YYYY-MM
    teacher: User = Depends(require_teacher),
    db: Session = Depends(get_db)
):
    school_class = db.query(SchoolClass).filter(SchoolClass.id == class_id).first()
    if not school_class:
        raise HTTPException(status_code=404, detail="រកមិនឃើញថ្នាក់រៀននេះទេ")
        
    enrollments = db.query(Enrollment).filter(Enrollment.class_id == class_id).order_by(Enrollment.roll_no.asc()).all()
    
    query = db.query(Attendance).filter(Attendance.class_id == class_id)
    if month:
        query = query.filter(func.to_char(Attendance.date, 'YYYY-MM') == month)
        
    attendances = query.order_by(Attendance.date.asc()).all()
    distinct_dates = sorted(list(set(str(a.date) for a in attendances)))
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance Summary"
    
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center")
    
    headers = ["ល.រ (Roll)", "អត្តលេខ (Code)", "ឈ្មោះសិស្ស (Name)"] + distinct_dates + ["វត្តមាន (Present)", "អវត្តមាន (Absent)", "ច្បាប់ (Permission)"]
    ws.append(headers)
    
    for col_num, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
        
    for e in enrollments:
        s = e.student
        student_att = {str(a.date): a.status for a in attendances if a.student_id == s.id}
        
        p_count = sum(1 for status in student_att.values() if status == "present")
        a_count = sum(1 for status in student_att.values() if status == "absent")
        perm_count = sum(1 for status in student_att.values() if status == "permission")
        
        row = [e.roll_no, s.student_code or "", s.name]
        for d in distinct_dates:
            row.append(student_att.get(d, "-"))
        row.extend([p_count, a_count, perm_count])
        ws.append(row)
        
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
        
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    month_suffix = f"_{month}" if month else ""
    filename = f"attendance_{school_class.name.replace(' ', '_')}{month_suffix}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )
